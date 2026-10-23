"""
Builds the two submission documents as plain .docx files (black text, no colours) for Google Docs.

  python Docs/build_docx.py      ->  Docs/TechnicalDocument.docx, Docs/VideoScript.docx

Diagrams are rendered from the SVGs in Docs/TechnicalDocument.html into Docs/figures/.
"""
import os, re, subprocess, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS = pathlib.Path(__file__).resolve().parent
EDGE = next(p for p in [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"] if os.path.exists(p))
BLACK = RGBColor(0, 0, 0)


def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)
    for name, size in (("Normal", 11), ("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.color.rgb = BLACK
        if name != "Normal": st.font.bold = True
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    return doc


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p


def bullets(doc, items):
    for it in items:
        if isinstance(it, tuple):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            doc.add_paragraph(it, style="List Bullet")


def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            r = cell.paragraphs[0].add_run(txt)
            r.font.size = Pt(10)
            r.bold = i == 0
    doc.add_paragraph()


def figure(doc, path, caption):
    doc.add_picture(str(path), width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9.5)


def render_figures():
    html = (DOCS / "TechnicalDocument.html").read_text(encoding="utf-8")
    out = DOCS / "figures"; out.mkdir(exist_ok=True)
    paths = []
    for i, svg in enumerate(re.findall(r"<svg.*?</svg>", html, flags=re.S)):
        w = int(re.search(r'width="(\d+)"', svg).group(1)); h = int(re.search(r'height="(\d+)"', svg).group(1))
        src = out / f"figure{i + 1}.html"
        src.write_text(f'<!doctype html><html><body style="margin:0;background:#fff">{svg}</body></html>', encoding="utf-8")
        png = out / f"figure{i + 1}.png"
        subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=2",
                        f"--window-size={w},{h}", f"--screenshot={png}", src.as_uri()], check=True, capture_output=True)
        src.unlink()
        paths.append(png)
    return paths


# ---------------------------------------------------------------- technical document
def technical(figs):
    doc = new_doc()
    doc.add_paragraph("Pre-Flight Drone Inspection Trainer", style="Title")
    para(doc, "Summative Assignment 2 · Technical Documentation · Unity 6 · Android")
    para(doc, "https://github.com/Strixx-thierry/AR-Pre-Flight-Drone-Inspection-Trainer", "GitHub: ")
    para(doc, "Android APK: [link]", "Mobile build: ")
    para(doc, "[link]", "Video: ")

    doc.add_heading("1. Overview", 1)
    para(doc, "The app trains drone pilots to do a pre-flight inspection. The trainee picks a drone (SKYFOLD F4 or TALON X4), "
              "chooses Guided or Assessment mode, and inspects a 3D drone: the battery latch, four propellers, the gimbal clamp "
              "and controller calibration. Faults are random each run. Unsafe launches are blocked, and every session ends with "
              "a scored report that is saved on the device.")
    para(doc, "Splash → Home → Hangar → Briefing → Inspection → Launch → Report → Records.", "Flow: ")
    para(doc, "The core simulation runs in a 3D viewer. AR is an optional add-on: on ARCore phones the Briefing screen "
              "also offers START IN AR, which places the same drone on a real surface.", "Modes: ")

    doc.add_heading("2. Architecture", 1)
    para(doc, "The code is split into a Runtime assembly (the game) and an Editor assembly (tools that generate scenes and assets). "
              "Each runtime subsystem has one job, and they talk through C# events: input goes to the task logic, and the UI reacts.")
    figure(doc, figs[0], "Figure 1: Runtime subsystems and event flow")
    bullets(doc, [
        ("Placement / view: ", "puts the drone in the scene and controls the camera (orbit, pan, zoom, focus)."),
        ("Drone: ", "DroneController and the inspectable parts; part state and tap selection."),
        ("TaskManager: ", "step order, rules, launch gate, errors, timer and score."),
        ("Directors: ", "MenuDirector and InspectionDirector control the screen flow and simulation states."),
        ("UI: ", "UIManager, 17 screens built from one UIFactory and theme."),
        ("Data and assets: ", "JSON scenarios, saved sessions, audio and Addressables loading."),
    ])

    doc.add_heading("3. AR add-on (AR Foundation)", 1)
    bullets(doc, [
        ("Session and planes: ", "ARSession with an XR Origin. ARPlaneManager detects horizontal planes, shown as a dotted surface."),
        ("Raycast placement: ", "ARRaycastManager raycasts from the screen. A ghost drone follows the surface and a tap places it at real size."),
        ("Adjust and confirm: ", "rotate, scale and re-place, then Confirm adds an ARAnchor so the drone stays fixed in the room."),
        ("Feedback: ", "the scan screen shows the surface count. If tracking is lost, an overlay pauses the timer until tracking returns or the user re-places the drone."),
        ("Design: ", "ARPlacementManager and the 3D viewer both implement IPlacementProvider, so the inspection logic is identical in both modes."),
    ])

    doc.add_heading("4. Simulation", 1)
    table(doc, [
        ["Entity", "Interaction", "Rule"],
        ["Battery latch", "Tap, slide to lock, swap", "Safe only when locked and charge is 2 LEDs or more"],
        ["4 propellers", "Tap, inspect, Airworthy / Damaged, replace", "Passing a damaged blade is a safety error"],
        ["Gimbal clamp", "Swipe up", "Must be removed before launch"],
        ["Controller", "Two virtual sticks", "All 8 positions must be reached"],
    ])
    para(doc, "In Guided mode the current part glows and later steps are locked. Launch Check refuses to launch while anything is "
              "unsafe, adds a safety error and points to the problem. When everything passes, the motors spin up, the drone lifts "
              "off, and the report shows time, errors, a step breakdown and a score.")
    para(doc, "Each run randomly picks 0–2 cracked blades, an unlatched or low battery, and a fitted or removed clamp. The weighting "
              "follows the UK CAA (CAP 2534), which names pre-flight checks and battery management as the main human-error causes.")

    doc.add_heading("5. OOP", 1)
    figure(doc, figs[1], "Figure 2: Inspectable part hierarchy")
    bullets(doc, [
        ("Encapsulation: ", "part state (locked, fault, fitted) is private and changed only through methods such as Toggle(), SubmitVerdict() and Remove()."),
        ("Abstraction: ", "AInspectableComponent defines any inspectable part. IPlacementProvider hides whether the 3D viewer or AR is used."),
        ("Inheritance: ", "RotorModule, BatteryLatchModule and GimbalCoverModule extend the base class. The 17 screens extend UIView."),
        ("Polymorphism: ", "the drone and the task logic loop over all parts and call IsSafe and ApplyScenario(). Each part applies its own rule."),
    ])

    doc.add_heading("6. Design patterns and events", 1)
    bullets(doc, [
        ("Observer: ", "parts raise StateChanged, and TaskManager raises Changed, Feedback and ErrorLogged. The HUD, audio and pop-ups "
                       "subscribe, so the logic never calls the UI directly."),
        ("State: ", "Placing, Inspection and Evaluation states, each with its own enter, update and exit."),
        ("Factory: ", "UIFactory and ChecklistRow.Create() build the checklist and UI from the JSON steps."),
        ("Object Pool: ", "FeedbackPopupPool uses Unity's ObjectPool for the floating feedback labels. Six are created once, then "
                          "reused through Get and Release instead of being created and destroyed for every action."),
        ("Singleton: ", "AudioManager, one instance that survives scene changes."),
    ])

    doc.add_heading("7. Data persistence", 1)
    bullets(doc, [
        "Scenario steps, hints and fault chances are JSON files in StreamingAssets, so they change without recompiling.",
        "Timer, current step and errors are kept in memory during a run.",
        "Each finished session (drone, mode, time, score, pass or fail, errors) is saved as JSON in PlayerPrefs, up to 20 sessions. "
        "The Records screen loads them at launch and shows the latest runs, best time and pass rate.",
    ])

    doc.add_heading("8. Audio", 1)
    para(doc, "AudioManager plays named sound effects from an AudioLibrary (13 clips) and has a separate loop for the motors. "
              "Sounds match events: a latch click, an error buzz, an alarm when launch is blocked, calibration ticks, and a motor "
              "spin-up at launch. Sound and vibration can be turned off in Settings.")

    doc.add_heading("9. Asset management", 1)
    para(doc, "The 3D drone models are in an Addressables group called Drones. A drone's definition stores only its address. When "
              "training starts, DroneAssetLoader calls Addressables.LoadAssetAsync for the chosen drone. When the scene closes, it "
              "destroys the instance and calls Addressables.Release, so only one model is in memory at a time. The menus use 2D "
              "pictures instead of the models.")

    doc.add_heading("10. Performance", 1)
    bullets(doc, [
        "URP mobile renderer at 0.8 render scale with one shadowed light.",
        "Light outline shader tinted with property blocks, with no extra materials.",
        "About 20k triangles per drone. 9-sliced UI sprites, pooled pop-ups and on-demand loading.",
    ])

    doc.add_heading("11. Changes from the SRS", 1)
    bullets(doc, [
        ("3D viewer as the core, AR as an add-on. ", "Agreed with the lecturer: inspecting small parts such as blade tips works better with free "
         "camera control, so the main simulation is 3D. AR placement is kept as an optional mode (START IN AR) so it can still be demonstrated."),
        ("Two drones instead of one, ", "showing that the checklist is data-driven."),
        ("Random faults, including low battery, ", "so trainees can't memorise the answers."),
        ("FR-08 corrected: ", "launch is prevented, not allowed, while critical items are unsafe."),
    ])

    doc.add_heading("12. Limitations and future work", 1)
    bullets(doc, [
        "AR mode needs an ARCore phone and good lighting. In AR the camera can't orbit underneath the drone.",
        "Saved data is local and single-user. A cloud store would allow instructor reports.",
        "More fault types (motors, compass, GPS) and a short flight test after the inspection.",
    ])
    doc.save(DOCS / "TechnicalDocument.docx")
    print("ok TechnicalDocument.docx")


# ---------------------------------------------------------------- video script
def script():
    doc = new_doc()
    doc.add_paragraph("Video Script", style="Title")
    para(doc, "About 8 minutes. Part 1 on the phone, Part 2 in Unity and the code.")

    doc.add_heading("Part 1: App on the phone (about 4 min)", 1)
    steps = [
        ("Open the app (Unity splash). ", "\"This is the Pre-Flight Drone Inspection Trainer on Android.\""),
        ("Home. ", "\"The menu has settings, help, records and the drone selector.\""),
        ("Help. ", "\"These are the instructions: rotate, zoom, tap parts, then launch check.\""),
        ("Hangar. ", "\"There are two drones, each with its own checklist.\" Select one."),
        ("Briefing. ", "\"Guided mode shows hints. Assessment mode scores everything.\" Press Start."),
        ("Rotate, zoom and pan the drone. ", "\"It's a full 3D model. I can look from any side.\""),
        ("Press Launch Check early. ", "\"Launch is blocked because the drone isn't safe. That's a safety error.\""),
        ("Battery. ", "\"Faults are random. I check the charge, swap it if it's low, and lock the latch.\""),
        ("Propellers. ", "\"The camera zooms to each blade. This one has a broken tip, so I mark it damaged and replace it.\""),
        ("Gimbal clamp. ", "\"I swipe the clamp off.\""),
        ("Calibrate RC. ", "\"I move both sticks to every position.\""),
        ("Launch. ", "\"Everything is safe now, so the drone takes off.\""),
        ("Report. ", "\"This shows the result, time, errors and score.\""),
        ("Records, then close and reopen the app. ", "\"Sessions are saved and still there after a restart.\""),
        ("AR add-on: Briefing ▸ START IN AR. ", "\"AR is an add-on. The app detects the floor or desk.\" Show the dotted planes."),
        ("Tap to place, rotate, confirm. ", "\"I place the drone at real size and confirm, which anchors it to the room. I can walk around it.\""),
        ("Tap the battery in AR. ", "\"The same inspection works in AR.\" Cover the camera briefly to show the tracking-lost overlay."),
    ]
    for i, (do, say) in enumerate(steps, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(do).bold = True
        p.add_run(say)

    doc.add_heading("Part 2: Unity and code (about 4 min)", 1)
    code = [
        ("Project folders. ", "\"The code is split into Runtime and Editor.\""),
        ("AInspectableComponent and RotorModule. ", "\"This is the base class for every part. Each part has its own safety rule.\""),
        ("TaskManager. ", "\"Parts send events, and the task manager checks the steps and blocks unsafe launches.\""),
        ("InspectionDirector. ", "\"The simulation moves through states: placing, inspecting, evaluating.\""),
        ("UIFactory. ", "\"The checklist is built from the JSON file.\""),
        ("FeedbackPopupPool. ", "\"The feedback labels are reused from an object pool.\""),
        ("Addressables Groups window and DroneAssetLoader. ", "\"Drones are loaded when needed and released afterwards.\""),
        ("Scenario JSON and ProgressStore. ", "\"Settings are JSON, and sessions are saved in PlayerPrefs.\""),
        ("AudioManager. ", "\"One audio manager plays all the sound effects.\""),
        ("ARPlacementManager. ", "\"This is the AR add-on: plane detection, raycast placement and anchors, behind the same interface as the 3D viewer.\""),
        ("DesignChanges document. ", "\"The 3D core and AR add-on decision is documented.\""),
    ]
    for i, (do, say) in enumerate(code, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(do).bold = True
        p.add_run(say)

    doc.add_heading("Before recording", 1)
    bullets(doc, ["Rebuild scenes and check the Console is clean.",
                  "Install the latest APK and do one full run so Records isn't empty.",
                  "Turn on Do Not Disturb, sound on, brightness up.",
                  "Open the code files in the order above."])
    doc.save(DOCS / "VideoScript.docx")
    print("ok VideoScript.docx")


if __name__ == "__main__":
    figs = render_figures()
    technical(figs)
    script()
